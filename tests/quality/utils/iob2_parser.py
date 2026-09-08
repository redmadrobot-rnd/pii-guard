"""Парсер IOB2-датасета для entity-level quality gate."""

from __future__ import annotations

import ast
import csv
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]

# Соответствие label'ов API и золотых IOB2-тегов из датасета.
API_TO_GOLD: dict[str, set[str]] = {
    "PERSON": {"FIRST_NAME", "LAST_NAME", "MIDDLE_NAME"},
    "LOCATION": {"CITY", "STREET", "REGION", "DISTRICT", "HOUSE", "COUNTRY"},
    "DATE_TIME": {"DATE_TIME"},
    "EMAIL_ADDRESS": {"EMAIL"},
    "PHONE_NUMBER": {"PHONE"},
    "PASSPORT": {"PASSPORT"},
    "DRIVER_LICENSE": {"DRIVER_LICENSE"},
    "MILITARY_ID": {"MILITARY_ID"},
    "BIRTH_CERTIFICATE": {"BIRTH_CERTIFICATE"},
    "OMS": {"OMS"},
    "SNILS": {"SNILS"},
    "INN": {"INN"},
    "CREDIT_CARD": {"CREDIT_CARD"},
    "BANK_ACCOUNT": {"BANK_ACCOUNT"},
    "POSTAL_CODE": {"POSTAL_CODE"},
    "TELEGRAM": {"TELEGRAM"},
    "BIK": {"BIK"},
    "IP_ADDRESS": {"IP_ADDRESS"},
    "IP_PORT": {"IP_ADDRESS"},
    "URL": {"URL"},
}


def _build_gold_to_api() -> dict[str, set[str]]:
    mapping: dict[str, set[str]] = defaultdict(set)
    for api_label, gold_labels in API_TO_GOLD.items():
        for gold_label in gold_labels:
            mapping[gold_label].add(api_label)
    return dict(mapping)


GOLD_TO_API: dict[str, set[str]] = _build_gold_to_api()


@dataclass
class GoldEntity:
    gold_label: str
    api_label: str
    start: int
    end: int
    text: str


@dataclass
class DatasetCase:
    row_id: str
    text: str
    source: str
    gold_entities: list[GoldEntity]


def _parse_list_cell(raw: str | None) -> list[str]:
    if not raw:
        return []
    try:
        value = ast.literal_eval(raw)
    except (ValueError, SyntaxError):
        return []
    if not isinstance(value, list):
        return []
    return [str(item) for item in value]


def _extract_case_text(row: dict[str, str], tokens: list[str]) -> str:
    text = str(row.get("text", "")).strip()
    if text:
        return text
    return " ".join(str(token) for token in tokens).strip()


def _find_token_spans(text: str, tokens: list[str]) -> list[tuple[int, int]]:
    """Ищет координаты токенов последовательно слева направо.

    Если токен не найден на текущей позиции, возвращается `(-1, -1)` без
    fallback-поиска от начала строки, чтобы не ломать монотонность координат.
    """
    spans: list[tuple[int, int]] = []
    pos = 0
    for token in tokens:
        idx = text.find(token, pos)
        if idx == -1:
            spans.append((-1, -1))
        else:
            spans.append((idx, idx + len(token)))
            pos = idx + len(token)
    return spans


def _extract_gold_entities(
    text: str,
    tokens: list[str],
    ner_tags: list[str],
) -> list[tuple[str, int, int, str]]:
    token_spans = _find_token_spans(text, tokens)
    entities: list[tuple[str, int, int, str]] = []
    cur_label: str | None = None
    cur_start: int | None = None
    cur_end: int | None = None

    for idx in range(min(len(tokens), len(ner_tags))):
        tag = str(ner_tags[idx])
        if tag == "O":
            label = None
        elif "-" in tag:
            label = tag.split("-", 1)[1]
        else:
            label = tag

        start, end = token_spans[idx]
        is_new = label != cur_label or tag.startswith("B-")

        if is_new:
            if cur_label is not None and cur_start is not None and cur_end is not None:
                entities.append(
                    (cur_label, cur_start, cur_end, text[cur_start:cur_end])
                )
            cur_label = label
            cur_start = start if start != -1 else None
            cur_end = end if end != -1 else None
        else:
            if end != -1 and cur_end is not None:
                cur_end = end

    if cur_label is not None and cur_start is not None and cur_end is not None:
        entities.append((cur_label, cur_start, cur_end, text[cur_start:cur_end]))

    return entities


def load_dataset_cases(
    dataset_path: Path,
) -> tuple[list[DatasetCase], set[str], set[str]]:
    """Читает IOB2 CSV и возвращает кейсы, покрытые категории и unmapped-лейблы."""
    if not dataset_path.exists():
        return [], set(), set()

    cases: list[DatasetCase] = []
    observed_api_labels: set[str] = set()
    unmapped_gold_labels: set[str] = set()

    for idx, row in enumerate(_iter_dataset_rows(dataset_path), start=1):
        tokens = _parse_list_cell(row.get("tokens"))
        ner_tags = _parse_list_cell(row.get("ner_tags"))
        text = _extract_case_text(row, tokens)
        source = str(row.get("source", "")).strip()

        if not text:
            continue

        gold_entities_raw = _extract_gold_entities(text, tokens, ner_tags)
        gold_entities: list[GoldEntity] = []
        for gold_label, start, end, entity_text in gold_entities_raw:
            api_candidates = GOLD_TO_API.get(gold_label, set())
            if not api_candidates:
                unmapped_gold_labels.add(gold_label)
                continue

            api_label = sorted(api_candidates)[0]
            observed_api_labels.add(api_label)
            gold_entities.append(
                GoldEntity(
                    gold_label=gold_label,
                    api_label=api_label,
                    start=start,
                    end=end,
                    text=entity_text,
                )
            )

        cases.append(
            DatasetCase(
                row_id=str(row.get("id", idx)).strip() or str(idx),
                text=text,
                source=source,
                gold_entities=gold_entities,
            )
        )

    return cases, observed_api_labels, unmapped_gold_labels


def _iter_dataset_rows(dataset_path: Path) -> list[dict[str, str]]:
    suffix = dataset_path.suffix.lower()
    if suffix == ".xlsx":
        return _iter_xlsx_rows(dataset_path)
    return _iter_csv_rows(dataset_path)


def _iter_csv_rows(dataset_path: Path) -> list[dict[str, str]]:
    with dataset_path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        return [dict(row) for row in reader]


def _iter_xlsx_rows(dataset_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(dataset_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return []
        columns = [str(item).strip() if item is not None else "" for item in header]
        out: list[dict[str, str]] = []
        for row in rows:
            payload: dict[str, str] = {}
            for idx, column in enumerate(columns):
                if not column:
                    continue
                value = row[idx] if row is not None and idx < len(row) else ""
                payload[column] = "" if value is None else str(value)
            out.append(payload)
        return out
    finally:
        workbook.close()
